from datetime import date
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROWS = 1000          # formula/validation capacity on entry sheets
SLOTS = 30           # products on the dashboard
CHART_SLOTS = 15     # products shown in the charts
FIRST = 9            # first product row on the dashboard
LAST = FIRST + SLOTS - 1

head_fill = PatternFill("solid", fgColor="1F4E78")
head_font = Font(bold=True, color="FFFFFF")
input_fill = PatternFill("solid", fgColor="FFF9E5")
thin = Side(style="thin", color="D0D0D0")
box = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()
pur, sal, dash = wb.active, wb.create_sheet("Sales"), wb.create_sheet("Dashboard")
pur.title = "Purchases"


def entry_sheet(ws, headers, samples):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.fill, cell.font, cell.border = head_fill, head_font, box
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for r in range(2, ROWS + 2):
        for c in range(1, 5):
            ws.cell(r, c).fill = input_fill
            ws.cell(r, c).border = box
        ws.cell(r, 5).border = box
        ws.cell(r, 2).number_format = "yyyy-mm-dd"
        ws.cell(r, 4).number_format = "#,##0.00"
        ws.cell(r, 5).value = f'=IF(OR(C{r}="",D{r}=""),"",C{r}*D{r})'
        ws.cell(r, 5).number_format = "#,##0.00"
    for i, row in enumerate(samples, 2):
        for c, v in enumerate(row, 1):
            ws.cell(i, c, v)
    for col, w in zip("ABCDE", (28, 14, 12, 16, 16)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{ROWS + 1}"
    dv = DataValidation(type="list", formula1=f"=Dashboard!$A${FIRST}:$A${LAST}",
                        allow_blank=True, showErrorMessage=True,
                        errorTitle="Unknown product",
                        error="Add the product to the Dashboard product list first.")
    ws.add_data_validation(dv)
    dv.add(f"A2:A{ROWS + 1}")
    dvd = DataValidation(type="date", operator="greaterThan", formula1="36526",
                         allow_blank=True, showErrorMessage=True,
                         error="Enter a valid date.")
    ws.add_data_validation(dvd)
    dvd.add(f"B2:B{ROWS + 1}")
    dvq = DataValidation(type="whole", operator="greaterThan", formula1="0",
                         allow_blank=True, showErrorMessage=True,
                         error="Quantity must be a whole number above 0.")
    ws.add_data_validation(dvq)
    dvq.add(f"C2:C{ROWS + 1}")


entry_sheet(pur, ["Product", "Date of Purchase", "Quantity", "Purchase Price (per unit)", "Total Cost"], [
    ("Widget", date(2026, 9, 1), 100, 4.50),
    ("Gadget", date(2026, 9, 3), 50, 12.00),
    ("Gizmo", date(2026, 9, 5), 200, 1.75),
    ("Widget", date(2026, 9, 20), 60, 4.80),
])
entry_sheet(sal, ["Product", "Date Sold", "Quantity Sold", "Sold Price (per unit)", "Total Revenue"], [
    ("Widget", date(2026, 9, 10), 70, 7.00),
    ("Gadget", date(2026, 9, 12), 20, 18.50),
    ("Gizmo", date(2026, 9, 15), 120, 3.25),
    ("Widget", date(2026, 9, 25), 40, 7.25),
])

# ---- Dashboard ----
from openpyxl.formatting.rule import CellIsRule

PERIODS = ["All time", "Today", "Yesterday", "Last 7 days", "Last 30 days",
           "This month", "Last month", "This year", "Custom range"]
dash["A1"] = "Inventory Dashboard"
dash["A1"].font = Font(bold=True, size=18, color="1F4E78")
dash["A2"] = ("Type product names in the yellow column (A). Everything else is calculated "
              "from the Purchases and Sales sheets. Names must match exactly.")
dash["A2"].font = Font(italic=True, color="666666")

# period list (helper, column N)
dash["N2"] = "Period options"
dash["N2"].font = Font(italic=True, color="999999")
for i, p in enumerate(PERIODS, 3):
    dash.cell(i, 14, p).font = Font(color="999999")
plist = f"$N$3:$N${2 + len(PERIODS)}"

dash["A3"], dash["B3"] = "Period", "All time"
dash["A4"], dash["C4"] = "Custom from", "Custom to"
dash["A5"], dash["C5"] = "Showing from", "to"
for ref in ("A3", "A4", "C4", "A5", "C5"):
    dash[ref].font = Font(bold=True, color="666666")
for ref in ("B3", "B4", "D4"):
    dash[ref].fill, dash[ref].border = input_fill, box
dash["B3"].font = Font(bold=True)
for ref in ("B4", "D4", "B5", "D5"):
    dash[ref].number_format = "yyyy-mm-dd"
dv = DataValidation(type="list", formula1=f"={plist}", allow_blank=False)
dash.add_data_validation(dv); dv.add("B3")
dvd = DataValidation(type="date", operator="greaterThan", formula1="36526",
                     allow_blank=True, showErrorMessage=True, error="Enter a valid date.")
dash.add_data_validation(dvd); dvd.add("B4"); dvd.add("D4")
dash["E5"] = "Days in period"
dash["E5"].font = Font(bold=True, color="666666")
dash["F5"] = (f'=MAX(1,MIN($D$5,TODAY())-MAX($B$5,MIN(Purchases!$B$2:$B${ROWS + 1},'
              f'Sales!$B$2:$B${ROWS + 1}))+1)')
dash["F5"].number_format = "0"
dash["G5"] = "(Days Stock Lasts = stock on hand / average units sold per day in this period)"
dash["G5"].font = Font(italic=True, color="888888")
dash["E4"] = "(used when Period = Custom range)"
dash["E4"].font = Font(italic=True, color="888888")

T = "TODAY()"
ym = f"YEAR({T}),MONTH({T})"
dash["B5"] = (f'=CHOOSE(MATCH($B$3,{plist},0),0,{T},{T}-1,{T}-6,{T}-29,'
              f'DATE({ym},1),DATE(YEAR({T}),MONTH({T})-1,1),DATE(YEAR({T}),1,1),IF($B$4="",0,$B$4))')
dash["D5"] = (f'=CHOOSE(MATCH($B$3,{plist},0),2958465,{T},{T}-1,{T},{T},'
              f'{T},EOMONTH({T},-1),{T},IF($D$4="",2958465,$D$4))')

KPI = 7
dash.cell(KPI, 1, "Total Cost"); dash.cell(KPI, 3, "Total Revenue")
dash.cell(KPI, 5, "Total Profit"); dash.cell(KPI, 7, "Overall Margin")
for c in (1, 3, 5, 7):
    dash.cell(KPI, c).font = Font(bold=True, color="666666")
kpi = {(KPI, 2): f"=SUM(F{FIRST}:F{LAST})", (KPI, 4): f"=SUM(G{FIRST}:G{LAST})",
       (KPI, 6): f"=SUM(J{FIRST}:J{LAST})", (KPI, 8): f'=IF(D{KPI}=0,"",F{KPI}/D{KPI})'}
for (r, c), f in kpi.items():
    cell = dash.cell(r, c, f)
    cell.font = Font(bold=True, size=13)
    cell.number_format = "0.0%" if c == 8 else "#,##0.00"

HDR = FIRST - 1
heads = ["Product", "Units Purchased", "Units Sold", "Stock On Hand (all time)",
         "Avg Buy Price (all time)", "Total Cost", "Total Revenue", "Avg Sell Price",
         "Cost of Goods Sold", "Profit", "Margin %", "Days Stock Lasts"]
for c, h in enumerate(heads, 1):
    cell = dash.cell(HDR, c, h)
    cell.fill, cell.font, cell.border = head_fill, head_font, box
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
dash.row_dimensions[HDR].height = 45

P, S = "Purchases", "Sales"
rng = lambda sh, col: f"{sh}!${col}$2:${col}${ROWS + 1}"
per = lambda sh, dcol, a: (f'{rng(sh,"B")},">="&$B$5,{rng(sh,"B")},"<="&$D$5')
for r in range(FIRST, LAST + 1):
    a = f"$A{r}"
    f = {
        2: f'SUMIFS({rng(P,"C")},{rng(P,"A")},{a},{per(P,0,0)})',
        3: f'SUMIFS({rng(S,"C")},{rng(S,"A")},{a},{per(S,0,0)})',
        4: f'SUMIF({rng(P,"A")},{a},{rng(P,"C")})-SUMIF({rng(S,"A")},{a},{rng(S,"C")})',
        5: f'IF(SUMIF({rng(P,"A")},{a},{rng(P,"C")})=0,0,SUMIF({rng(P,"A")},{a},{rng(P,"E")})/SUMIF({rng(P,"A")},{a},{rng(P,"C")}))',
        6: f'SUMIFS({rng(P,"E")},{rng(P,"A")},{a},{per(P,0,0)})',
        7: f'SUMIFS({rng(S,"E")},{rng(S,"A")},{a},{per(S,0,0)})',
        8: f'IF(C{r}=0,0,G{r}/C{r})',
        9: f'C{r}*E{r}',
        10: f'G{r}-I{r}',
        11: f'IF(G{r}=0,0,J{r}/G{r})',
        12: f'IF(C{r}=0,"No sales",IF(D{r}<=0,0,ROUND(D{r}/(C{r}/$F$5),1)))',
    }
    dash.cell(r, 1).fill = input_fill
    for c in range(1, 13):
        dash.cell(r, c).border = box
        if c > 1:
            dash.cell(r, c, f'=IF({a}="","",{f[c]})')
            dash.cell(r, c).number_format = "0.0%" if c == 11 else ("0.0" if c == 12 else ("#,##0" if c in (2, 3, 4) else "#,##0.00"))
            if c == 12:
                dash.cell(r, c).alignment = Alignment(horizontal="right")
for r, name in enumerate(("Widget", "Gadget", "Gizmo"), FIRST):
    dash.cell(r, 1, name)
dash.column_dimensions["A"].width = 26
for col in "BCDEFGHIJKL":
    dash.column_dimensions[col].width = 14
dash.column_dimensions["N"].width = 16
dash.freeze_panes = f"B{FIRST}"

red = Font(color="C00000", bold=True)
dash.conditional_formatting.add(f"D{FIRST}:D{LAST}", CellIsRule(operator="lessThan", formula=["0"], font=red))
dash.conditional_formatting.add(f"L{FIRST}:L{LAST}", CellIsRule(operator="lessThan", formula=["7"], font=red))
dash.conditional_formatting.add(f"J{FIRST}:K{LAST}", CellIsRule(operator="lessThan", formula=["0"], font=red))

last_c = FIRST + CHART_SLOTS - 1
cats = Reference(dash, min_col=1, min_row=FIRST, max_row=last_c)


def chart(title, cols, anchor, ytitle):
    ch = BarChart()
    ch.type, ch.title, ch.height, ch.width = "col", title, 8.5, 16
    for col in cols:
        ch.add_data(Reference(dash, min_col=col, min_row=HDR, max_row=last_c), titles_from_data=True)
    ch.set_categories(cats)
    ch.y_axis.title = ytitle
    ch.y_axis.delete = ch.x_axis.delete = False
    ch.legend.position = "b"
    dash.add_chart(ch, anchor)


chart("Units Purchased vs Sold", (2, 3), f"A{LAST + 3}", "Units")
chart("Stock On Hand", (4,), f"F{LAST + 3}", "Units")
chart("Cost vs Revenue", (6, 7), f"A{LAST + 21}", "Amount")
chart("Profit by Product", (10,), f"F{LAST + 21}", "Profit")

wb.active = 2
wb.save("inventory_tracker.xlsx")
print("saved")
